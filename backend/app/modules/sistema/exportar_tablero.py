"""RF-GER-14 a 18: el Excel del tablero del gerente.

El gerente pidio dos cosas en la junta y la segunda se olvida facil: ver los
cuatro indicadores de las unidades entre dos fechas del calendario, Y poder
bajarlos a Excel. Lo segundo no es un capricho -- el reporte que el sube a
direccion se arma en Excel, y si la aplicacion no se lo da, lo que va a hacer
es teclear a mano los numeros que ve en la grafica. Ahi empieza el problema que
este archivo existe para evitar.

DE DONDE SALEN LOS NUMEROS, y por que de un solo lado:

  TODO sale de estadisticas.tablero_rango(). Ni un indicador se recalcula aqui.
  Esa regla no es de estilo: es la misma leccion que ya esta escrita en
  exportar_resumen.py del taller. El area tiene hoy una hoja Graficos que no
  cuadra con su hoja RESUMEN porque los cuatro numeros se copian de una a otra
  a mano, y el dia que alguien cambio un rango solo cambio una. Cuando la
  pantalla y el archivo discrepan, el gerente deja de creerle a los dos.

  Lo unico que este modulo consulta por su cuenta son las hojas de DETALLE
  (CITAS y ORDENES), que no son indicadores sino el renglon por renglon de
  donde salieron. Para decidir a que periodo pertenece cada fila se reusa
  estadisticas._clave() y la lista de etiquetas que ya devolvio tablero_rango,
  en vez de volver a calcular donde empieza y termina cada periodo: esa cuenta
  vive en _periodos_entre() y tenerla en dos lados es justo como se llega a que
  el total y el detalle digan cosas distintas.

EL FORMATO ES EL DEL OTRO EXCEL A PROPOSITO. Los colores, las fuentes, los
encabezados con relleno y los paneles fijos se IMPORTAN de
taller/exportar_resumen.py, no se copian. Son dos archivos que la misma gente
va a tener abiertos al mismo tiempo, y si se copiaran las constantes bastaria
con que alguien cambiara un azul en un lado para que dejaran de parecer del
mismo sistema.

QUE TOTALES VAN COMO FORMULA Y CUALES NO, que es la decision menos obvia de
este archivo:

  citas_totales y atenciones SI van como formula (=SUM sobre la hoja SERIES).
  Son conteos: el total del rango es exactamente la suma de las barras, asi que
  la formula no vuelve a definir el indicador --suma los numeros que ya salieron
  de tablero_rango-- y de paso deja el archivo cuadrando solo, que es el punto
  del Excel del area.

  dias_reparacion_prom y tasa_concretadas NO pueden ir como formula sobre esas
  columnas, y conviene dejar escrito por que antes de que alguien lo "arregle":
  el promedio del rango NO es el promedio de los promedios por periodo. Un mes
  con una sola orden de 30 dias pesaria igual que uno con cuarenta ordenes de un
  dia. Lo mismo con el porcentaje: un =AVERAGE de la columna de por cientos da
  un numero distinto del que sale de dividir los totales. Van como valor, tal
  como los calculo tablero_rango, que si acumula suma y conteo por separado.
"""
import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy.orm import Session

from ... import models as m
from ...core.tiempo import dia_operativo
from . import estadisticas
# Los estilos se importan del Excel del taller en vez de copiarse. Ver la nota
# de arriba: son los dos archivos que salen de la misma aplicacion y tienen que
# verse iguales sin que nadie se acuerde de sincronizarlos.
from ..taller.exportar_resumen import (VERDE, _anchos, _borde,  # noqa: E402
                                       _encabezados, _etiqueta, _hoja_titulo,
                                       _relleno_sec)

# Donde arranca la primera fila de datos de una hoja con encabezado, contando
# como lo hacen _hoja_titulo() (titulo en la 1, renglon en blanco en la 2) y
# _encabezados() (encabezado en la 3). Se deja como constante porque la hoja
# RESUMEN necesita apuntar con formulas a filas concretas de SERIES, y una
# formula que apunte una fila mas arriba de lo debido no truena: suma un numero
# distinto y nadie lo nota.
PRIMERA_FILA = 4

_nota = Font(name="Calibri", size=9, italic=True, color="6B7A80")
_aviso = Font(name="Calibri", size=10, bold=True, color="8A4B08")
_numero = Font(name="Calibri", size=11, bold=True, color=VERDE)

_GRANULARIDADES = {"dia": "Diaria", "mes": "Mensual", "anio": "Anual"}


def _renglon(ws, fila: int, etiqueta: str, valor, formato: str | None = None,
             fuente: Font | None = None) -> int:
    """Un renglon de 'etiqueta | valor' de los bloques de la hoja RESUMEN."""
    ws.cell(row=fila, column=1, value=etiqueta).border = _borde
    c = ws.cell(row=fila, column=2, value=valor)
    c.border = _borde
    c.alignment = Alignment(horizontal="right")
    if formato:
        c.number_format = formato
    if fuente:
        c.font = fuente
    return fila + 1


def _bloque(ws, fila: int, titulo: str) -> int:
    """El encabezado gris de una seccion, igual que en el Excel del taller."""
    c = ws.cell(row=fila, column=1, value=titulo)
    c.font = _etiqueta
    c.fill = _relleno_sec
    ws.cell(row=fila, column=2).fill = _relleno_sec
    return fila + 1


def _hoja_resumen(wb, tablero: dict, hoy: datetime.date):
    ws = wb.create_sheet("RESUMEN")
    _anchos(ws, [40, 20, 14])
    fila = _hoja_titulo(ws, "Tablero del gerente", 3)

    desde = datetime.date.fromisoformat(tablero["desde"])
    hasta = datetime.date.fromisoformat(tablero["hasta"])
    res = tablero["resumen"]

    fila = _bloque(ws, fila, "RANGO CONSULTADO")
    fila = _renglon(ws, fila, "Desde", desde, "DD/MM/YYYY")
    fila = _renglon(ws, fila, "Hasta", hasta, "DD/MM/YYYY")
    fila = _renglon(ws, fila, "Granularidad",
                    _GRANULARIDADES.get(tablero["granularidad"], tablero["granularidad"]))
    fila = _renglon(ws, fila, "Periodos en el rango", len(tablero["etiquetas"]))
    # El periodo en curso se dice con todas sus letras en vez de dejarlo como
    # una etiqueta suelta: quien abra el archivo tres meses despues tiene que
    # poder saber si la ultima barra estaba a medias cuando se exporto, porque
    # un mes incompleto contra meses completos siempre se ve como una caida.
    fila = _renglon(ws, fila, "Periodo sin terminar al exportar",
                    tablero["en_curso"] or "ninguno (el rango termina en el pasado)")
    fila = _renglon(ws, fila, "Archivo generado el", hoy, "DD/MM/YYYY")
    fila += 1

    # La hoja SERIES se arma despues, pero sus filas son predecibles: una por
    # etiqueta a partir de PRIMERA_FILA. Se calculan aqui para que los dos
    # conteos del resumen apunten a ellas.
    ultima = PRIMERA_FILA + len(tablero["etiquetas"]) - 1

    fila = _bloque(ws, fila, "INDICADORES DEL PERIODO COMPLETO")
    # El None se escribe como celda VACIA y no como cero. Es la misma decision
    # que ya esta tomada en cumplimiento_choferes(): un 0.0 dias de reparacion
    # promedio afirma que se reparo en el acto, y un 0% de citas concretadas
    # pinta de incumplido a quien no fallo a nada. El hueco tiene que verse como
    # hueco, sobre todo aqui, donde alguien va a graficar la columna despues.
    fila = _renglon(ws, fila, "Tiempo promedio de reparacion (dias)",
                    res["dias_reparacion_prom"], "0.0", _numero)
    fila = _renglon(ws, fila, "Citas concretadas (%)",
                    res["tasa_concretadas"], "0.0", _numero)
    fila = _renglon(ws, fila, "Citas totales",
                    "=SUM(SERIES!D%d:D%d)" % (PRIMERA_FILA, ultima), "0", _numero)
    fila = _renglon(ws, fila, "Atenciones a citas",
                    "=SUM(SERIES!E%d:E%d)" % (PRIMERA_FILA, ultima), "0", _numero)
    fila = _renglon(ws, fila, "Ordenes cerradas en el rango",
                    res["ordenes_cerradas"], "0")
    fila += 1

    if tablero.get("aviso"):
        # Si el rango se recorto por el tope de periodos hay que decirlo TAMBIEN
        # en el archivo. En la pantalla el aviso se ve una vez y se cierra; el
        # Excel se reenvia por correo y se lee meses despues, y ahi la grafica
        # empieza donde el gerente no pidio sin nada que lo explique.
        c = ws.cell(row=fila, column=1, value=tablero["aviso"])
        c.font = _aviso
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 1, end_column=3)
        fila += 3

    nota = ws.cell(row=fila, column=1, value=(
        "Citas totales y Atenciones son formula: suman la hoja SERIES y se "
        "recalculan solas. Los otros dos van como valor a proposito -- el "
        "promedio del rango no es el promedio de los promedios por periodo, y "
        "el porcentaje del rango no es el promedio de los porcentajes. Una "
        "celda vacia significa que no hubo base para calcular ese periodo, que "
        "no es lo mismo que cero."))
    nota.font = _nota
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 4, end_column=3)


def _hoja_series(wb, tablero: dict):
    ws = wb.create_sheet("SERIES")
    _anchos(ws, [16, 26, 20, 16, 20])
    fila = _hoja_titulo(ws, "Indicadores periodo por periodo", 5)
    fila = _encabezados(ws, fila, [
        "PERIODO", "TIEMPO PROM. DE REPARACION (DIAS)", "CITAS CONCRETADAS (%)",
        "CITAS TOTALES", "ATENCIONES A CITAS"])

    # Los cuatro vectores se toman por clave y no por posicion del arreglo
    # `series`: el contrato fija el orden, pero leerlo por clave es lo que hace
    # que este archivo siga saliendo bien si algun dia se agrega un quinto
    # indicador en medio.
    datos = {s["clave"]: s["datos"] for s in tablero["series"]}
    for i, etq in enumerate(tablero["etiquetas"]):
        # None se pasa tal cual: openpyxl deja la celda VACIA, que es lo que
        # tiene que quedar cuando no hubo ordenes cerradas o no hubo citas con
        # desenlace. Escribir 0 seria inventar una medicion.
        ws.append([etq,
                   datos["dias_reparacion"][i], datos["tasa_concretadas"][i],
                   datos["citas_totales"][i], datos["atenciones"][i]])

    for f in ws.iter_rows(min_row=fila, max_row=ws.max_row):
        for c in f:
            c.border = _borde
        f[1].number_format = "0.0"
        f[2].number_format = "0.0"
        if f[0].value == tablero["en_curso"]:
            # El periodo que todavia no termina se sombrea en vez de omitirse.
            # Omitirlo dejaria la grafica sin el dato de hoy; dejarlo sin marca
            # lo hace parecer un desplome. La nota de abajo dice que significa
            # el sombreado, porque un color sin leyenda no informa nada.
            for c in f:
                c.fill = _relleno_sec

    fila = ws.max_row + 2
    nota = ws.cell(row=fila, column=1, value=(
        "El renglon sombreado es el periodo que todavia no terminaba cuando se "
        "exporto el archivo: va a la mitad y no se compara contra los "
        "completos. Una celda vacia no es un cero, es que no hubo con que "
        "calcular -- sin ordenes cerradas no hay promedio de reparacion, y sin "
        "citas cumplidas ni faltas no hay porcentaje que sacar."))
    nota.font = _nota
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 3, end_column=5)


def _hoja_citas(wb, db: Session, tablero: dict):
    ws = wb.create_sheet("CITAS")
    _anchos(ws, [14, 14, 16, 28])
    fila = _hoja_titulo(ws, "Citas del rango", 4)
    fila = _encabezados(ws, fila, ["UNIDAD", "FECHA", "ESTADO", "TALLER"])

    gran = tablero["granularidad"]
    vivos = set(tablero["etiquetas"])
    citas = [c for c in db.query(m.CitaTaller).all()
             if estadisticas._clave(c.fecha_cita, gran) in vivos]
    for c in sorted(citas, key=lambda x: x.fecha_cita):
        ws.append([c.unidad.num_economico if c.unidad else None,
                   c.fecha_cita, c.estado,
                   c.taller.nombre if c.taller else None])

    ws.auto_filter.ref = "A%d:D%d" % (fila - 1, max(ws.max_row, fila))
    for f in ws.iter_rows(min_row=fila, max_row=ws.max_row):
        for c in f:
            c.border = _borde
        f[1].number_format = "DD/MM/YYYY"

    fila = ws.max_row + 2
    nota = ws.cell(row=fila, column=1, value=(
        "Aqui aparecen TODAS las citas del rango, canceladas incluidas, pero "
        "el indicador Citas totales no las cuenta: una cita cancelada nunca "
        "llego a ser un compromiso vivo. O sea que el numero de renglones de "
        "esta hoja es mayor que el total de la hoja RESUMEN, y la diferencia "
        "son exactamente las canceladas. Se listan porque son la respuesta a "
        "la pregunta que sigue cuando el total baja."))
    nota.font = _nota
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 3, end_column=4)


def _hoja_ordenes(wb, db: Session, tablero: dict):
    ws = wb.create_sheet("ORDENES")
    _anchos(ws, [18, 14, 14, 14, 10, 18])
    fila = _hoja_titulo(ws, "Ordenes cerradas en el rango", 6)
    fila = _encabezados(ws, fila,
                        ["FOLIO", "UNIDAD", "ENTRADA", "SALIDA", "DIAS", "TIPO"])

    gran = tablero["granularidad"]
    vivos = set(tablero["etiquetas"])
    ordenes = []
    for o in (db.query(m.OrdenServicio)
              .filter(m.OrdenServicio.fecha_salida.isnot(None)).all()):
        # El periodo de la orden se saca de su DIA OPERATIVO en Tijuana, igual
        # que en tablero_rango(). fecha_salida es UTCDateTime: una orden cerrada
        # a las 17:30 del 30 de septiembre en el taller esta guardada como las
        # 00:30 del 1 de octubre UTC, y tomarle la fecha cruda la mandaria al
        # mes siguiente. Aqui eso no seria un decimal de mas: seria una orden
        # que aparece en el detalle de un mes en el que no se cerro.
        salida = dia_operativo(o.fecha_salida)
        if estadisticas._clave(salida, gran) not in vivos:
            continue
        entrada = dia_operativo(o.fecha_entrada) if o.fecha_entrada else None
        dias = (o.fecha_salida - o.fecha_entrada).days if o.fecha_entrada else None
        if dias is not None and dias < 0:
            dias = None
        ordenes.append((o, entrada, salida, dias))

    for o, entrada, salida, dias in sorted(ordenes, key=lambda x: x[2]):
        # DIAS va vacio --no en cero-- cuando la orden no tiene entrada o salio
        # antes de entrar. Son las mismas dos que tablero_rango() deja fuera del
        # promedio, y dejarlas aqui en blanco hace que el detalle se pueda leer
        # contra el indicador sin tener que explicar de donde salio la
        # diferencia: las que no tienen numero son las que no se promediaron.
        ws.append([o.folio, o.unidad.num_economico if o.unidad else None,
                   entrada, salida, dias, o.tipo])

    ws.auto_filter.ref = "A%d:F%d" % (fila - 1, max(ws.max_row, fila))
    for f in ws.iter_rows(min_row=fila, max_row=ws.max_row):
        for c in f:
            c.border = _borde
        f[2].number_format = "DD/MM/YYYY"
        f[3].number_format = "DD/MM/YYYY"

    fila = ws.max_row + 2
    nota = ws.cell(row=fila, column=1, value=(
        "La orden se cuenta en el periodo de su SALIDA, no en el de su "
        "entrada: una unidad que entro en enero y salio en marzo habla de "
        "marzo, que es cuando por fin libero el espacio. DIAS son periodos "
        "completos de 24 horas entre entrada y salida, truncados hacia abajo: "
        "una unidad que entro por la tarde y salio en la manana de dos dias "
        "despues cuenta 1, no 2. Un DIAS vacio es una orden sin fecha de "
        "entrada o con las fechas al reves, y esas quedan fuera del promedio."))
    nota.font = _nota
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 4, end_column=6)


def construir(db: Session, desde: datetime.date, hasta: datetime.date,
              gran: str = "mes",
              hoy: datetime.date | None = None) -> tuple[bytes, str]:
    """Devuelve (.xlsx listo para descargar, nombre del archivo).

    Devuelve tambien el nombre en vez de dejar que lo arme el controlador
    porque el rango del nombre tiene que ser el que REALMENTE se exporto: si el
    gerente invirtio los calendarios o si el tope de periodos recorto el rango,
    tablero_rango() corrige las fechas, y un archivo que se llama distinto de lo
    que trae adentro es el que termina citado en una junta con el rango
    equivocado.
    """
    hoy = hoy or datetime.date.today()
    # Una sola llamada, una sola fuente. Todo lo que se escriba en las cuatro
    # hojas sale de aqui o del detalle que apunta a estas mismas etiquetas.
    tablero = estadisticas.tablero_rango(db, desde, hasta, gran)

    wb = Workbook()
    wb.remove(wb.active)
    _hoja_resumen(wb, tablero, hoy)
    _hoja_series(wb, tablero)
    _hoja_citas(wb, db, tablero)
    _hoja_ordenes(wb, db, tablero)

    nombre = "Tablero gerente %s a %s %s.xlsx" % (
        tablero["desde"], tablero["hasta"],
        _GRANULARIDADES.get(tablero["granularidad"], tablero["granularidad"]).lower())

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), nombre
