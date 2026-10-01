"""Los Excel de las areas entran por la pantalla, ya no por scp.

POR QUE EXISTE. El motor que lee los Excel (importadores/sincronizar.py) corre
cada madrugada contra la carpeta de las areas, pero para que un archivo nuevo
llegara a esa carpeta alguien tenia que copiarlo al servidor por scp y meterlo al
contenedor. Ningun departamento podia hacerlo solo, asi que los catalogos
envejecian. Esto es la puerta de entrada:

  1. El area sube su Excel. Se revisa que sea un Excel de verdad y se guarda
     aparte, sin tocar la carpeta de las areas.
  2. Se prueba con el MISMO motor de la madrugada, sobre una COPIA de la base y
     con una copia de la carpeta donde el archivo nuevo ocupa el lugar del viejo.
     Se le ensena a la persona que cambiaria, fila por fila.
  3. Si confirma, el viejo se guarda en historial/, el nuevo toma su lugar y se
     sincroniza de verdad (sincronizar respalda la base antes). Queda en la
     bitacora de auditoria quien subio que.

Cada area sube solo lo suyo; el administrador, todo. El servidor lo valida aqui
aunque la pantalla ya esconda lo que no le toca a cada quien.

Los textos que lee la persona (titulo de cada archivo, que trae) viven en la
pantalla: aqui solo las claves y los nombres de archivo.
"""
import datetime
import glob
import hashlib
import io
import json
import os
import secrets
import shutil
import tempfile
import threading

from ... import importador
from ...core.tiempo import TZ_OPERACION, ahora_utc
from ...importadores import flota, personal, sincronizar, taller

AREAS = {"logistica": "Logistica", "almacen": "Almacen", "compras": "Compras",
         "taller": "Taller"}
ROL_DE_AREA = {a: "datos_" + a for a in AREAS}
ROLES_DE_CARGA = ("administrador", *ROL_DE_AREA.values())

# tipo -> de que area es, con que nombre lo busca el importador y que pasos de
# sincronizar lo leen. `patron`: el libro REQUIS trae la version en el nombre y
# el importador lo busca por pedazo, no por nombre exacto.
TIPOS = {
    "unidades": {"area": "logistica", "archivo": flota.ARCHIVO,
                 "pasos": ("flota", "catalogo")},
    "choferes": {"area": "logistica", "archivo": personal.ARCHIVO_INFO,
                 "pasos": ("base", "personal")},
    "choferes_lan": {"area": "logistica", "archivo": personal.ARCHIVO_LAN,
                     "pasos": ("personal",)},
    "requisiciones": {"area": "compras", "patron": importador.REQUIS_PATRON,
                      "pasos": ("base",)},
    "codigos": {"area": "almacen", "archivo": "CODIGOS TALLER.xlsx",
                "pasos": ("base",)},
    "resumen": {"area": "taller", "archivo": taller.ARCHIVO, "pasos": ("taller",)},
}

EXTENSIONES = (".xlsx", ".xlsm")
TOPE_BYTES = 25 * 1024 * 1024
HISTORIAL_POR_TIPO = 15
DIAS_EN_ESPERA = 2

# Umbrales del "revisa bien antes de aplicar". Un archivo equivocado en el lugar
# del catalogo de unidades (otra hoja, otro libro) puede dar de baja la flota
# entera, y el simulacro lo diria con un numero que nadie lee.
CAMBIADAS_MUCHAS = 50

# Una sola carga a la vez en este proceso: dos sincronizaciones encimadas
# respaldan y escriben la misma base al mismo tiempo.
_candado = threading.Lock()


class ErrorCarga(Exception):
    """Un problema que la persona puede entender y corregir."""

    def __init__(self, mensaje: str, status: int = 422):
        super().__init__(mensaje)
        self.status = status


def carpeta_areas() -> str:
    """Donde lee el reloj de la madrugada. En el servidor, /datos/areas."""
    return (os.environ.get("CARPETA_AREAS")
            or ("/datos/areas" if os.path.isdir("/datos") else importador.CARPETA_DATOS))


def _carpeta_espera(areas: str) -> str:
    return os.path.join(areas, ".cargas")


def tipos_permitidos(roles) -> list:
    roles = set(roles)
    if "administrador" in roles:
        return list(TIPOS)
    areas = {a for a, rol in ROL_DE_AREA.items() if rol in roles}
    return [t for t, d in TIPOS.items() if d["area"] in areas]


def _es_de(tipo: str, nombre: str) -> bool:
    """Si este archivo de la carpeta es el que ocupa el lugar de `tipo`."""
    if nombre.startswith("~$"):           # archivo de bloqueo de Excel
        return False
    d = TIPOS[tipo]
    if "patron" in d:
        return d["patron"] in nombre.upper() and nombre.lower().endswith(EXTENSIONES)
    return nombre == d["archivo"]


def _nombre_destino(tipo: str, original: str) -> str:
    d = TIPOS[tipo]
    if "patron" in d:
        ext = os.path.splitext(original)[1].lower()
        return d["patron"] + (ext if ext in EXTENSIONES else ".xlsx")
    return d["archivo"]


def _hora(ts: float) -> str:
    return datetime.datetime.fromtimestamp(ts, TZ_OPERACION).isoformat(timespec="seconds")


def estado(areas: str, roles) -> dict:
    """Lo que la pantalla ensena antes de subir nada."""
    archivos = []
    nombres = sorted(os.listdir(areas)) if os.path.isdir(areas) else []
    for tipo in tipos_permitidos(roles):
        actual = next((n for n in nombres if _es_de(tipo, n)), None)
        info = None
        if actual:
            st = os.stat(os.path.join(areas, actual))
            info = {"nombre": actual, "actualizado": _hora(st.st_mtime), "bytes": st.st_size}
        archivos.append({"tipo": tipo, "area": TIPOS[tipo]["area"], "actual": info})
    ultima = None
    ruta = os.path.join(areas, "ultima-sincronizacion.json")
    if os.path.exists(ruta):
        try:
            with open(ruta, encoding="utf-8") as f:
                u = json.load(f)
            ultima = {"fin": u.get("fin"),
                      "pasos": [{"paso": p["paso"], "estado": p["estado"]}
                                for p in u.get("pasos", [])]}
        except (OSError, ValueError):
            ultima = None
    return {"archivos": archivos, "ultima": ultima}


def validar_excel(nombre: str, contenido: bytes):
    """Que sea un libro de Excel que se pueda abrir, antes de gastar un simulacro."""
    if not nombre.lower().endswith(EXTENSIONES):
        raise ErrorCarga("Solo se aceptan libros de Excel (.xlsx o .xlsm). "
                         "Si lo tienes en .xls o .csv, abrelo en Excel y guardalo como .xlsx.")
    if len(contenido) > TOPE_BYTES:
        raise ErrorCarga("El archivo pesa mas de 25 MB.", 413)
    if not contenido.startswith(b"PK"):
        raise ErrorCarga("El archivo no es un libro de Excel valido.")
    from openpyxl import load_workbook
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True)
        hojas = libro.sheetnames
        libro.close()
    except Exception:
        raise ErrorCarga("Excel no pudo abrir este archivo: puede estar danado o protegido.")
    return hojas


def _limpiar_espera(espera: str):
    if not os.path.isdir(espera):
        return
    limite = ahora_utc().timestamp() - DIAS_EN_ESPERA * 86400
    for n in os.listdir(espera):
        ruta = os.path.join(espera, n)
        if os.path.isdir(ruta) and os.path.getmtime(ruta) < limite:
            shutil.rmtree(ruta, ignore_errors=True)


def _alertas(cambios: dict) -> list:
    out = []
    for tabla, x in sorted(cambios.items()):
        if x["borradas"]:
            out.append({"tabla": tabla, "tipo": "borradas", "cuantas": x["borradas"]})
        if x["cambiadas"] >= CAMBIADAS_MUCHAS:
            out.append({"tabla": tabla, "tipo": "cambiadas", "cuantas": x["cambiadas"]})
    return out


def _veredicto(tipo: str, informe: dict):
    """(puede_aplicar, motivo). Solo cuentan los pasos que leen ESTE archivo:
    que otra area tenga roto el suyo no es razon para frenar a esta."""
    for p in informe["pasos"]:
        if p["paso"] in TIPOS[tipo]["pasos"] and p["estado"] != "ok":
            return False, "%s: %s" % (p["paso"], p.get("motivo") or p["estado"])
    return True, None


def _resumen(tipo: str, informe: dict) -> dict:
    puede, motivo = _veredicto(tipo, informe)
    return {"pasos": informe["pasos"], "pasos_propios": list(TIPOS[tipo]["pasos"]),
            "cambios": informe["cambios"],
            "ejemplos": informe.get("ejemplos", {}),
            "cuentas_nuevas": len(informe.get("cuentas_nuevas", [])),
            "alertas": _alertas(informe["cambios"]),
            "puede_aplicar": puede, "motivo": motivo}


def simular(areas: str, tipo: str, nombre_original: str, contenido: bytes,
            usuario_id: int) -> dict:
    """Guarda el archivo en espera y dice que cambiaria. No toca nada real."""
    validar_excel(nombre_original, contenido)
    espera = _carpeta_espera(areas)
    os.makedirs(espera, exist_ok=True)
    _limpiar_espera(espera)

    ahora = ahora_utc()
    carga_id = ahora.astimezone(TZ_OPERACION).strftime("%Y%m%d-%H%M%S-") + secrets.token_hex(3)
    dir_carga = os.path.join(espera, carga_id)
    os.makedirs(dir_carga)
    destino = os.path.join(dir_carga, _nombre_destino(tipo, nombre_original))
    with open(destino, "wb") as f:
        f.write(contenido)

    # La carpeta de prueba: todo lo de las areas tal cual, salvo el lugar de
    # este tipo, que lo ocupa el archivo nuevo. Los demas pasos corren con sus
    # archivos de siempre y, como son idempotentes, no cambian nada: lo que
    # aparece en el simulacro es lo que trae ESTE archivo.
    prueba = tempfile.mkdtemp(prefix="carga-")
    try:
        if os.path.isdir(areas):
            for n in os.listdir(areas):
                p = os.path.join(areas, n)
                if (os.path.isfile(p) and n.endswith(EXTENSIONES) and not n.startswith("~$")
                        and not _es_de(tipo, n)):
                    shutil.copy2(p, os.path.join(prueba, n))
        shutil.copy2(destino, os.path.join(prueba, os.path.basename(destino)))
        informe = sincronizar.sincronizar(prueba, simular=True)
    finally:
        shutil.rmtree(prueba, ignore_errors=True)

    resumen = _resumen(tipo, informe)
    meta = {"carga_id": carga_id, "tipo": tipo, "original": nombre_original,
            "archivo": os.path.basename(destino),
            "sha256": hashlib.sha256(contenido).hexdigest(), "bytes": len(contenido),
            "usuario_id": usuario_id, "simulado_en": ahora.timestamp(),
            "puede_aplicar": resumen["puede_aplicar"], "aplicada": None}
    with open(os.path.join(dir_carga, "carga.json"), "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False)
    return {"carga_id": carga_id, "tipo": tipo, "archivo": meta["archivo"],
            "original": nombre_original, **resumen}


def _podar_historial(historial: str, tipo: str):
    viejos = sorted(glob.glob(os.path.join(glob.escape(historial), "*--%s--*" % tipo)))
    for v in viejos[:-HISTORIAL_POR_TIPO]:
        os.remove(v)


def aplicar(areas: str, carga_id: str, roles) -> dict:
    """Pone el archivo en su lugar y sincroniza de verdad. Devuelve el informe
    y lo que hay que anotar en la bitacora (lo anota quien llama, con su sesion)."""
    if not carga_id.replace("-", "").isalnum():
        raise ErrorCarga("Carga inexistente.", 404)
    dir_carga = os.path.join(_carpeta_espera(areas), carga_id)
    ruta_meta = os.path.join(dir_carga, "carga.json")
    if not os.path.exists(ruta_meta):
        raise ErrorCarga("Esta revision ya no existe (se borran a los dos dias). "
                         "Vuelve a subir el archivo.", 404)
    with open(ruta_meta, encoding="utf-8") as f:
        meta = json.load(f)
    tipo = meta["tipo"]
    if tipo not in tipos_permitidos(roles):
        raise ErrorCarga("Este archivo no es de tu area.", 403)
    if meta.get("aplicada"):
        raise ErrorCarga("Este archivo ya se aplico.", 409)
    if not meta.get("puede_aplicar"):
        raise ErrorCarga("La revision encontro un error con este archivo: no se puede aplicar.")

    # Si algo sincronizo DESPUES de la revision (el reloj de la madrugada u otra
    # carga), lo que se le enseno a la persona ya no es lo que pasaria.
    ultima = os.path.join(areas, "ultima-sincronizacion.json")
    if os.path.exists(ultima) and os.path.getmtime(ultima) > meta["simulado_en"]:
        raise ErrorCarga("Los datos cambiaron desde que revisaste este archivo. "
                         "Vuelve a subirlo para ver los cambios al dia.", 409)

    if not _candado.acquire(blocking=False):
        raise ErrorCarga("Hay otra carga aplicandose en este momento. Intenta en un minuto.", 409)
    try:
        historial = os.path.join(areas, "historial")
        os.makedirs(historial, exist_ok=True)
        sello = ahora_utc().astimezone(TZ_OPERACION).strftime("%Y%m%d-%H%M%S")
        movidos = []
        for n in os.listdir(areas):
            p = os.path.join(areas, n)
            if os.path.isfile(p) and _es_de(tipo, n):
                guardado = os.path.join(historial, "%s--%s--%s" % (sello, tipo, n))
                shutil.move(p, guardado)
                movidos.append((guardado, p))
        nuevo = os.path.join(areas, meta["archivo"])
        shutil.copy2(os.path.join(dir_carga, meta["archivo"]), nuevo)

        informe = sincronizar.sincronizar(areas, simular=False)
        puede, motivo = _veredicto(tipo, informe)
        if not puede:
            # El paso que lee este archivo ya revirtio lo suyo; se regresa
            # tambien el archivo, o la madrugada volveria a tropezar con el.
            os.remove(nuevo)
            for guardado, original in movidos:
                shutil.move(guardado, original)
            raise ErrorCarga("No se pudo aplicar: %s. Se dejo el archivo anterior." % motivo)
        _podar_historial(historial, tipo)
        meta["aplicada"] = ahora_utc().isoformat()
        with open(ruta_meta, "w", encoding="utf-8") as f:
            json.dump(meta, f, ensure_ascii=False)
    finally:
        _candado.release()

    resumen = _resumen(tipo, informe)
    asiento = {"tipo": tipo, "archivo": meta["archivo"], "original": meta["original"],
               "sha256": meta["sha256"], "respaldo": informe.get("respaldo"),
               "cambios": {t: {k: v for k, v in x.items() if not k.startswith("_")}
                           for t, x in informe["cambios"].items()}}
    return {"tipo": tipo, "archivo": meta["archivo"], "respaldo": informe.get("respaldo"),
            **resumen, "asiento": asiento}
