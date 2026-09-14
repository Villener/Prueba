"""Arma el Excel del taller con el formato que el area ya conoce.

El administrador lleva anos entregando RESUMEN.xlsx, asi que el archivo que
salga de aqui tiene que parecerse a ese o nadie lo va a usar. Se conservan sus
secciones y sus etiquetas exactas; lo que cambia es de donde salen los numeros.

TRES DIFERENCIAS CON EL ARCHIVO DEL AREA, Y LAS TRES SON A PROPOSITO:

1. LOS TOTALES SON FORMULA, NO NUMEROS TECLEADOS. En el archivo del area solo
   3 de sus 27 filas son formula y el resto se cuenta a mano contra PATIO cada
   dia. De ahi salen sus errores: al 27 de agosto tenian 3 y 4 en las casillas
   de refacciones chinas y talleres externos, cuando su propia hoja PATIO dice
   4 y 3. El total cuadraba, asi que nadie lo noto.

2. SE AGREGA LA CUBETA "MAS DE 366 DIAS". La del area corta en 366 y mete todo
   lo mas viejo en esa ultima casilla. Hoy hay una unidad --la 1017-- con mas
   de 1200 dias parada, y contarla junto a una de un ano esconde justo el caso
   que habria que estar mirando.

3. NO SE INVENTA LA SERIE HISTORICA DIARIA. El area tiene 204 columnas, una por
   dia. Eso NO se puede reconstruir desde la base, y la razon importa: su hoja
   REPARADO no anota la fecha de SALIDA en el 98% de los registros, asi que no
   hay forma de saber que habia en el patio un martes de marzo. Se entrega la
   foto de HOY, que si es exacta, y el conteo de ENTRADAS por mes, que tambien.
   Inventar la serie seria dar por cierto algo que nadie registro.
"""
import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ... import models as m

# Las cubetas viven en indicadores.py y NO se copian aqui: si el Excel tuviera
# su propia definicion, el dia que alguien cambie un rango la pantalla y el
# archivo dirian cosas distintas. Es el error que el area tiene hoy entre su
# hoja Graficos y su hoja RESUMEN.
from .indicadores import POR_ANTIGUEDAD, POR_ESTADO  # noqa: E402

AZUL = "0F3452"
VERDE = "0E6235"
GRIS = "EEF0EB"

_titulo = Font(name="Calibri", size=14, bold=True, color=AZUL)
_encabezado = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
_etiqueta = Font(name="Calibri", size=11, bold=True, color=AZUL)
_relleno_enc = PatternFill("solid", fgColor=AZUL)
_relleno_sec = PatternFill("solid", fgColor=GRIS)
_borde = Border(*[Side(style="thin", color="D5DAD2")] * 4)


def _hoja_titulo(ws, texto: str, ancho: int, fila: int = 1) -> int:
    ws.cell(row=fila, column=1, value=texto).font = _titulo
    return fila + 2


def _encabezados(ws, fila: int, nombres: list) -> int:
    for i, nombre in enumerate(nombres, start=1):
        c = ws.cell(row=fila, column=i, value=nombre)
        c.font = _encabezado
        c.fill = _relleno_enc
        c.alignment = Alignment(horizontal="left", vertical="center")
    # Se fija el panel por REFERENCIA de texto, no con ws.cell(). Pedirle la
    # celda a openpyxl la CREA, y con eso max_row sube una fila: el primer
    # append caia en la fila siguiente y el archivo salia con un renglon vacio
    # entre el encabezado y los datos.
    ws.freeze_panes = f"A{fila + 1}"
    return fila + 1


def _anchos(ws, anchos: list):
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a


def _dias(ingreso: datetime.date, hoy: datetime.date) -> int:
    return (hoy - ingreso).days


def _hoja_resumen(wb, adentro: list, entradas_por_mes: list, hoy: datetime.date):
    ws = wb.create_sheet("RESUMEN")
    _anchos(ws, [34, 14, 12])
    fila = _hoja_titulo(ws, f"Resumen del taller al {hoy:%d/%m/%Y}", 3)

    def bloque(nombre: str, renglones: list, fila: int) -> int:
        c = ws.cell(row=fila, column=1, value=nombre)
        c.font = _etiqueta
        c.fill = _relleno_sec
        ws.cell(row=fila, column=2).fill = _relleno_sec
        fila += 1
        primera = fila
        for etiqueta, cuantas in renglones:
            ws.cell(row=fila, column=1, value=etiqueta).border = _borde
            cc = ws.cell(row=fila, column=2, value=cuantas)
            cc.border = _borde
            cc.alignment = Alignment(horizontal="right")
            fila += 1
        ws.cell(row=fila, column=1, value="TOTAL").font = _etiqueta
        # Formula, no numero: es lo que impide el error de dedo del archivo viejo.
        t = ws.cell(row=fila, column=2,
                    value=f"=SUM(B{primera}:B{fila - 1})")
        t.font = _etiqueta
        t.alignment = Alignment(horizontal="right")
        return fila + 2

    por_estado = [(etq, sum(1 for x in adentro if x.clasificacion == cod))
                  for cod, etq in POR_ESTADO]
    sin_clasificar = sum(1 for x in adentro if x.clasificacion is None)
    if sin_clasificar:
        por_estado.append(("(sin clasificar)", sin_clasificar))
    fila = bloque("POR ESTADO", por_estado, fila)

    por_edad = []
    for etq, desde, hasta, _alerta in POR_ANTIGUEDAD:
        por_edad.append((etq, sum(1 for x in adentro
                                  if desde <= _dias(x.fecha_ingreso, hoy) <= hasta)))
    fila = bloque("POR ANTIGUEDAD EN EL PATIO", por_edad, fila)

    usos = {}
    for x in adentro:
        usos[x.uso or "(sin area)"] = usos.get(x.uso or "(sin area)", 0) + 1
    fila = bloque("POR AREA / USO", sorted(usos.items()), fila)

    # Entradas por mes: esto SI se puede calcular, porque la fecha de ingreso
    # existe en el 100% de los registros.
    c = ws.cell(row=fila, column=1, value="ENTRADAS POR MES (ultimos 12)")
    c.font = _etiqueta
    c.fill = _relleno_sec
    ws.cell(row=fila, column=2).fill = _relleno_sec
    fila += 1
    for mes, cuantas in entradas_por_mes:
        ws.cell(row=fila, column=1, value=mes).border = _borde
        cc = ws.cell(row=fila, column=2, value=cuantas)
        cc.border = _borde
        cc.alignment = Alignment(horizontal="right")
        fila += 1

    fila += 1
    nota = ws.cell(row=fila, column=1, value=(
        "Los totales son formula: se recalculan solos. La serie diaria del "
        "archivo anterior no se reproduce porque la fecha de salida solo existe "
        "en el 2% del historial, y sin ella no se puede saber que habia en el "
        "patio un dia pasado."))
    nota.font = Font(name="Calibri", size=9, italic=True, color="6B7A80")
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 3, end_column=3)


def _hoja_patio(wb, adentro: list, hoy: datetime.date):
    ws = wb.create_sheet("PATIO")
    _anchos(ws, [12, 8, 18, 40, 16, 14, 14, 18, 12, 8, 34, 6, 12])
    fila = _hoja_titulo(ws, f"Unidades en el taller al {hoy:%d/%m/%Y}", 13)
    fila = _encabezados(ws, fila, [
        "UNIDAD", "MOD", "MARCA", "FALLA", "MEC/SUPER", "USO", "STATUS",
        "OBSERVACIONES", "INGRESO", "DIAS", "PARTES SOLICITADAS", "L", "AREA"])
    for x in sorted(adentro, key=lambda a: a.fecha_ingreso):
        u = x.unidad
        ws.append([x.unidad_texto or (u.num_economico if u else ""),
                   u.anio if u else None, u.modelo if u else None,
                   x.falla, x.mecanico_texto, x.uso, x.estatus, x.observaciones,
                   x.fecha_ingreso, _dias(x.fecha_ingreso, hoy),
                   x.partes_solicitadas, x.clasificacion, x.area])
    for f in ws.iter_rows(min_row=fila, max_row=ws.max_row):
        for c in f:
            c.border = _borde
        f[8].number_format = "DD/MM/YYYY"


def _hoja_reparado(wb, historial: list):
    ws = wb.create_sheet("REPARADO")
    _anchos(ws, [12, 8, 18, 40, 16, 14, 14, 18, 12, 12, 34, 6, 12])
    fila = _hoja_titulo(ws, "Historial de entradas a taller", 13)
    # Una sola forma para toda la hoja. La del area cambia de layout dos veces
    # a media hoja, que es de donde salia el riesgo de leerla mal.
    fila = _encabezados(ws, fila, [
        "UNIDAD", "MOD", "MARCA", "FALLA", "MEC/SUPER", "USO", "STATUS",
        "OBSERVACIONES", "INGRESO", "SALIDA", "PARTES SOLICITADAS", "L", "AREA"])
    for x in historial:
        u = x.unidad
        ws.append([x.unidad_texto or (u.num_economico if u else ""),
                   u.anio if u else None, u.modelo if u else None,
                   x.falla, x.mecanico_texto, x.uso, x.estatus, x.observaciones,
                   x.fecha_ingreso, x.fecha_salida,
                   x.partes_solicitadas, x.clasificacion, x.area])
    ws.auto_filter.ref = f"A{fila - 1}:M{max(ws.max_row, fila)}"
    for f in ws.iter_rows(min_row=fila, max_row=ws.max_row):
        f[8].number_format = "DD/MM/YYYY"
        f[9].number_format = "DD/MM/YYYY"


def construir(db: Session, taller_nombre: str | None = None,
              hoy: datetime.date | None = None) -> bytes:
    """Devuelve el .xlsx listo para descargar."""
    hoy = hoy or datetime.date.today()

    q = db.query(m.MovimientoTaller).join(m.Unidad, isouter=True)
    if taller_nombre:
        q = q.filter(m.MovimientoTaller.area == taller_nombre.upper())

    adentro = [x for x in q.filter(m.MovimientoTaller.sigue_adentro.is_(True)).all()]
    historial = (q.filter(m.MovimientoTaller.sigue_adentro.is_(False))
                  .order_by(m.MovimientoTaller.fecha_ingreso.desc()).all())

    # Entradas por mes, de los ultimos 12 meses con datos.
    por_mes = {}
    for x in adentro + historial:
        clave = f"{x.fecha_ingreso:%Y-%m}"
        por_mes[clave] = por_mes.get(clave, 0) + 1
    entradas = sorted(por_mes.items())[-12:]

    wb = Workbook()
    wb.remove(wb.active)
    _hoja_resumen(wb, adentro, entradas, hoy)
    _hoja_patio(wb, adentro, hoy)
    _hoja_reparado(wb, historial)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
