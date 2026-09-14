"""Carga el historial del taller: las hojas REPARADO y PATIO de RESUMEN.xlsx.

Son 1,869 entradas a taller desde diciembre de 2021 sobre 536 unidades, mas las
que estan en el patio ahora mismo. Es de donde van a salir los graficos del
administrador y la exportacion con formato de RESUMEN, asi que si esto entra
mal, todo lo que se calcule despues sale mal y con cara de estar bien.

EL PROBLEMA DE ESTA HOJA, Y POR QUE NO SE PUEDE LEER POR POSICION.

REPARADO no tiene un layout: tiene TRES, y cambian a media hoja sin aviso.

  filas ~2-1000    la fecha de ingreso esta en la columna J y la K esta vacia
  filas ~1000-2065 la J pasa a ser observaciones y la fecha se va a la K
  filas 2067-2131  todo se recorre una columna: la G deja de ser basura y pasa
                   a ser el USO, la H el STATUS, la I las observaciones, la
                   fecha vuelve a la J -- y aparece fecha de SALIDA en la O

Un importador que lea 'la columna K' guarda texto en un campo de fecha para un
tercio de la hoja, y nadie se entera, porque el resultado sigue pareciendo
plausible. Aqui el layout se detecta POR CONTENIDO, fila por fila, anclando en
la columna de STATUS: solo puede tener tres valores, asi que encontrarla dice
donde esta todo lo demas.

La fecha se busca igual: se toma el primer valor que REALMENTE sea una fecha
entre las dos columnas candidatas, en vez de confiar en cual deberia ser.

Ademas, las filas donde solo hay una fecha en la columna A no son movimientos:
son separadores que dicen 'aqui empieza el dia tal'.

Se ejecuta con:  python -m app.importadores.taller
"""
import datetime
import logging
import os

import openpyxl
from sqlalchemy.orm import Session

from .. import models as m
from . import normaliza as n

log = logging.getLogger(__name__)

CARPETA_DATOS = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))))), "datos")

ARCHIVO = "RESUMEN.xlsx"

# Los tres unicos valores que toma STATUS. Son el ancla que dice que layout
# trae la fila.
ESTATUS = {"PORINGRESAR", "ENPROCESO", "PEDIDOCHINO"}

# El USO alimenta el tercer bloque de la hoja RESUMEN (filas 19-27).
USOS = {"REPARTO", "PIPA", "PIPAS", "UTILITARIA", "UTILITARIAS", "OPERACIONES",
        "LECHERIA", "AZUCENA", "FORANEA", "FORANEAS", "TUBERIA", "MASICOS"}


def _clave(v) -> str:
    return n.clave_unidad(v)


def _fecha(v):
    """La celda como fecha, o None. Los enteros (dias calculados) no son fecha."""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def _norm(v) -> str:
    return _clave(v)


def _leer_reparado(ws) -> list:
    """Una fila por movimiento. Detecta el layout por contenido, no por posicion."""
    salida, separadores, sin_fecha = [], 0, 0

    for fila in ws.iter_rows(min_row=2, values_only=True):
        def c(i):
            return fila[i] if i < len(fila) else None

        unidad = n.texto(c(1))          # B
        if not unidad:
            # Fila separadora: solo trae la fecha del dia en la columna A.
            if _fecha(c(0)) is not None:
                separadores += 1
            continue

        # ---- donde esta el STATUS decide donde esta todo lo demas -----------
        if _norm(c(7)) in ESTATUS:      # H
            uso, estatus, obs = c(6), c(7), c(8)          # G, H, I
        elif _norm(c(8)) in ESTATUS:    # I
            uso, estatus, obs = c(7), c(8), c(9)          # H, I, J
        else:
            # Sin status reconocible. Se rescata el USO si alguna de las dos
            # columnas candidatas trae uno, y se sigue: perder la fila entera
            # por no reconocer una palabra seria peor.
            uso = c(6) if _norm(c(6)) in USOS else (c(7) if _norm(c(7)) in USOS else None)
            estatus, obs = None, None

        # ---- la fecha: la que REALMENTE lo sea, no la que deberia -----------
        ingreso = _fecha(c(9)) or _fecha(c(10))           # J, luego K
        if ingreso is None:
            sin_fecha += 1
            continue

        # Las observaciones nunca son una fecha; si lo son, es que esa columna
        # era en realidad el ingreso y aqui no hay observaciones.
        if _fecha(obs) is not None:
            obs = None

        clasif = c(12)                                    # M
        clasif = int(clasif) if isinstance(clasif, (int, float)) and 1 <= clasif <= 5 else None

        salida.append({
            "unidad": unidad,
            "clave": _clave(unidad),
            "anio": n.anio(c(2)),
            "marca": n.texto(c(3)),
            "falla": n.texto(c(4)),
            "mecanico": n.texto(c(5)),
            "uso": n.texto(uso).upper() or None,
            "estatus": n.texto(estatus).upper() or None,
            "observaciones": n.texto(obs).upper() or None,
            "ingreso": ingreso,
            "salida": _fecha(c(14)),                       # O
            "partes": n.texto(c(11)) or None,              # L
            "clasificacion": clasif,
            "area": n.texto(c(13)).upper() or None,        # N
            "sigue_adentro": False,
            "origen": "REPARADO",
        })
    return salida, separadores, sin_fecha


def _leer_patio(ws) -> list:
    """PATIO si tiene encabezado de verdad (fila 2), asi que se mapea por nombre."""
    filas = list(ws.iter_rows(values_only=True))
    if len(filas) < 3:
        return []
    cab = {}
    for i, v in enumerate(filas[1]):
        clave = _norm(v)
        if clave and clave not in cab:
            cab[clave] = i

    def val(fila, *nombres):
        for nombre in nombres:
            i = cab.get(_norm(nombre))
            if i is not None and i < len(fila):
                return fila[i]
        return None

    salida = []
    for fila in filas[2:]:
        unidad = n.texto(val(fila, "UNIDAD"))
        if not unidad or _clave(unidad) == "":
            continue
        ingreso = _fecha(val(fila, "INGRESO"))
        if ingreso is None:
            continue
        clasif = val(fila, "L")
        clasif = int(clasif) if isinstance(clasif, (int, float)) and 1 <= clasif <= 5 else None
        salida.append({
            "unidad": unidad,
            "clave": _clave(unidad),
            "anio": n.anio(val(fila, "MOD")),
            "marca": n.texto(val(fila, "MARCA")),
            "falla": n.texto(val(fila, "FALLA")),
            "mecanico": n.texto(val(fila, "MEC/SUPER", "MECANICO/SUPERVISOR")),
            "uso": n.texto(val(fila, "USO")).upper() or None,
            "estatus": n.texto(val(fila, "STATUS")).upper() or None,
            # El encabezado esta mal escrito en el archivo del area. Se busca
            # como esta, no como deberia estar.
            "observaciones": n.texto(val(fila, "OBSEREVACIONES", "OBSERVACIONES")).upper() or None,
            "ingreso": ingreso,
            "salida": None,
            "partes": n.texto(val(fila, "PARTES SOLICITADA")) or None,
            "clasificacion": clasif,
            "area": n.texto(val(fila, "AREA")).upper() or None,
            "sigue_adentro": True,
            "origen": "PATIO",
        })
    return salida


def _tipo_por_uso(uso: str, tipos: dict) -> int:
    u = _norm(uso)
    if "PIPA" in u:
        return tipos.get("pipa", tipos["utilitario"])
    if "REPARTO" in u:
        return tipos.get("reparto", tipos["utilitario"])
    return tipos["utilitario"]


def importar(db: Session, carpeta: str = CARPETA_DATOS) -> dict:
    wb = openpyxl.load_workbook(os.path.join(carpeta, ARCHIVO),
                                read_only=True, data_only=True)
    reparado, separadores, sin_fecha = _leer_reparado(wb["REPARADO"])
    patio = _leer_patio(wb["PATIO"])
    wb.close()

    tipos = {t.nombre: t.id for t in db.query(m.TipoUnidad).all()}
    por_clave = {}
    for u in db.query(m.Unidad).all():
        c = _clave(u.num_economico)
        if c:
            por_clave.setdefault(c, u)

    ya = {(mv.unidad_id, mv.fecha_ingreso)
          for mv in db.query(m.MovimientoTaller.unidad_id,
                             m.MovimientoTaller.fecha_ingreso).all()}

    r = {"reparado": len(reparado), "patio": len(patio),
         "separadores": separadores, "sin_fecha": sin_fecha,
         "insertados": 0, "repetidos": 0, "unidades_creadas": [],
         "sin_clasificacion": 0}

    # PATIO primero: si una unidad esta en las dos hojas con la misma fecha,
    # manda la version que dice que SIGUE ADENTRO, que es la del presente.
    for fila in patio + reparado:
        u, _ = n.resolver_unidad(fila["clave"], por_clave)
        if u is None:
            # El taller la atiende pero Logistica nunca la dio de alta. Se crea
            # marcada, porque descartarla se lleva el 12% del historial.
            u = m.Unidad(
                num_economico=fila["unidad"][:20],
                tipo_unidad_id=_tipo_por_uso(fila["uso"] or "", tipos),
                estado="disponible", activo=True, origen="taller",
                anio=fila["anio"], modelo=(fila["marca"] or None),
            )
            db.add(u)
            db.flush()
            por_clave[fila["clave"]] = u
            r["unidades_creadas"].append(fila["unidad"])

        llave = (u.id, fila["ingreso"])
        if llave in ya:
            r["repetidos"] += 1
            continue
        ya.add(llave)

        if fila["clasificacion"] is None:
            r["sin_clasificacion"] += 1

        db.add(m.MovimientoTaller(
            unidad_id=u.id, unidad_texto=fila["unidad"][:40],
            fecha_ingreso=fila["ingreso"], fecha_salida=fila["salida"],
            sigue_adentro=fila["sigue_adentro"],
            falla=fila["falla"] or None, mecanico_texto=(fila["mecanico"] or None),
            uso=fila["uso"], estatus=fila["estatus"],
            observaciones=(fila["observaciones"] or None)[:80] if fila["observaciones"] else None,
            partes_solicitadas=fila["partes"], clasificacion=fila["clasificacion"],
            area=fila["area"], origen=fila["origen"],
        ))
        r["insertados"] += 1

    db.commit()
    return r


if __name__ == "__main__":  # pragma: no cover
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    from ..core.database import SessionLocal

    db = SessionLocal()
    try:
        antes = db.query(m.MovimientoTaller).count()
        r = importar(db)
        print(f"REPARADO leido        : {r['reparado']} movimientos "
              f"({r['separadores']} filas separadoras, {r['sin_fecha']} sin fecha)")
        print(f"PATIO leido           : {r['patio']} unidades adentro hoy")
        print(f"movimientos guardados : {antes} -> {db.query(m.MovimientoTaller).count()}")
        print(f"  insertados          : {r['insertados']}")
        print(f"  repetidos (omitidos): {r['repetidos']}")
        print(f"  sin clasificar 1-5  : {r['sin_clasificacion']}")
        creadas = r["unidades_creadas"]
        print(f"unidades que Logistica no tiene dadas de alta: {len(creadas)}")
        if creadas:
            print("  " + ", ".join(sorted(set(creadas))[:24])
                  + (" ..." if len(set(creadas)) > 24 else ""))
    finally:
        db.close()
