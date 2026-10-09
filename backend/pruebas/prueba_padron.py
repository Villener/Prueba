"""Pruebas del padron de unidades: el control de GPS de Logistica (2026-10-09).

    cd backend
    .venv/Scripts/python.exe pruebas/prueba_padron.py     (Windows)

Desde el 2026-10-09 el padron de unidades activas es el control de GPS y no
UNIDADES BAJA GAS. Cada caso arma un libro chico con el MISMO acomodo que el
real (FLOTILLA, GPS LISTADO, BASE DE DATOS y la tabla de permisos de
PORCENTAJE) y una base en memoria. Los nombres, VIN y placas son inventados: el
repositorio es publico.

- La planta: la que digan dos de las tres hojas. El canal, el de FLOTILLA.
- El permiso, del canal y la planta.
- BASE DE DATOS: las que esperan su GPS entran; las marcadas BAJA, no.
- Las del padron quedan activas con su planta y su permiso; las demas, de baja.
- No se duplica una unidad que la base guarda con otro nombre ('GRUA 718').
- Placas y VIN: solo los que faltan; lo distinto se reporta, no se pisa.
- Correrlo dos veces no cambia nada.
"""
import datetime
import os
import shutil
import sys
import tempfile
import traceback

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND)
os.environ["DATABASE_URL"] = "sqlite://"

import openpyxl                                             # noqa: E402
from sqlalchemy import create_engine                        # noqa: E402
from sqlalchemy.orm import sessionmaker                     # noqa: E402

from app import models as m                                 # noqa: E402
from app import seed                                        # noqa: E402
from app.core.database import Base                          # noqa: E402
from app.importadores import padron, sincronizar            # noqa: E402
from app.modules.bitacora import bitacora_controller as bc  # noqa: E402
from app.modules.bitacora import bitacora_service as bs     # noqa: E402
from app.modules.sistema import cargas_service as cs        # noqa: E402
from app.modules.taller import administrador_controller as ac  # noqa: E402

CASOS = []
HOY = datetime.date(2026, 10, 9)


def caso(nombre):
    def deco(fn):
        CASOS.append((nombre, fn))
        return fn
    return deco


P_ALAMOS, P_VALLE, P_TECATE = ("LP/14792/DIST/PLA/2016", "LP/14586/DIST/PLA/2016",
                               "LP/14791/DIST/PLA/2016")
# (canal, permiso, zona, supervisor), como la tabla de PORCENTAJE: el canal solo
# en el primer renglon de cada grupo.
PERMISOS = [
    ("REPARTO", P_ALAMOS, "ALAMOS", "ANA LOPEZ"),
    (None, P_VALLE, "ROSARITO", "BRUNO CASTRO"),
    (None, P_TECATE, "TECATE", "CARMEN DIAZ"),
    (None, P_VALLE, "VALLE REDONDO", "DANIEL ESPINO"),
    ("ESTACIONARIO", P_ALAMOS, "ALAMOS", "IRMA JUAREZ"),
    (None, P_ALAMOS, "ROSARITO", "BRUNO CASTRO"),
    (None, P_VALLE, "VALLE REDONDO", "ELENA FUENTES"),
    ("OPERACIÓN", P_VALLE, "VALLE REDONDO", "ING. FELIPE"),
    ("A- COMERCIAL", None, "ALAMOS", "GLORIA HERRERA"),
    ("TALLER", None, "ALAMOS", "JORGE KURI"),
    ("TESORERIA", None, "20 DE NOVIEMBRE", "HUGO IBARRA"),
]
# (canal, sucursal, supervisor, unidad, status). FLOTILLA es la hoja oculta y
# vieja: escribe al supervisor con su nombre completo.
FLOTILLA = [
    ("REPARTO", "ALAMOS", "ANA LOPEZ RUIZ", 2154, "LISTO"),
    ("REPARTO", "ALAMOS", "ANA LOPEZ RUIZ", 2155, "LISTO"),
    ("REPARTO", "ALAMOS", "ANA LOPEZ RUIZ", 2400, "LISTO"),       # GPS: otro supervisor
    ("REPARTO", "ROSARITO", "BRUNO CASTRO LEAL", 2310, "LISTO"),
    ("REPARTO", "ROSARITO", "BRUNO CASTRO LEAL", "BG-813", "LISTO"),  # GPS: estacionario
    ("ESTACIONARIO", "ROSARITO", "BRUNO CASTRO LEAL", "BG-733P", "LISTO"),
    ("ESTACIONARIO", "ROSARITO", "BRUNO CASTRO LEAL", 2311, "LISTO"),  # vieja: es de Irma
    ("ESTACIONARIO", "ALAMOS", "IRMA JUAREZ", "BG-740P", "LISTO"),
    ("ESTACIONARIO", "ALAMOS", "IRMA JUAREZ", "BG-741P", "LISTO"),
    ("ESTACIONARIO", "VALLE REDONDO", "ELENA FUENTES", "BG-728P", "PENDIENTE"),
    ("REPARTO", "VALLE REDONDO", "DANIEL ESPINO", "GB-304", "PENDIENTE"),
    ("REPARTO", "VALLE REDONDO", "DANIEL ESPINO", 1489, "PENDIENTE"),   # BAJA en BASE
    ("REPARTO", "TECATE", "CARMEN DIAZ", 1512, "PENDIENTE"),            # BAJA en la nota
    ("OPERACIONES", "VALLE REDONDO", "ING. FELIPE GOMEZ", "T-822", "LISTO"),
    ("OPERACIONES", "VALLE REDONDO", "ING. FELIPE GOMEZ", "BGOT-04", "PENDIENTE"),
    ("OPERACIONES", "VALLE REDONDO", "ING. FELIPE GOMEZ", "BGTO-05", "LISTO"),
    ("OPERACIONES", "VALLE REDONDO", "ING. FELIPE GOMEZ", "900R", "LISTO"),
]
# (unidad, modelo, vin, placas, canal, supervisor). El celular es el chip del GPS.
GPS = [
    (2154, "CNJ", "LS1TEST0000000001", "BG-0001-A", "REPARTO", "ANA LOPEZ"),
    (2400, "CNJ", None, None, "REPARTO", "DANIEL ESPINO"),
    (2310, "CNJ", "VINDEOTRAUNIDAD01", "BG0002A", "REPARTO", "BRUNO CASTRO"),
    ("BG-813", "CNJ", None, None, "ESTACIONARIO", "BRUNO CASTRO"),
    (2311, "CNJ", None, None, "ESTACIONARIO", "IRMA JUAREZ"),
    ("BG-728", "KENWORTH", "3BKTEST0000000003", None, "ESTACIONARIO", "ELENA FUENTES"),
    ("BG-304", "CNJ", None, None, "REPARTO", "DANIEL ESPINO"),
    ("GB-306", "CNJ", None, None, "REPARTO", "DANIEL ESPINO"),    # nueva, con la errata
    ("BG-834", "KENWORTH", None, None, "ESTACIONARIO", "BRUNO CASTRO"),
    (1167, "GM SPARK", None, None, "A. COMERCIAL", "GLORIA HERRERA"),
    (1203, "GM SUBURBAN", None, None, "TESORERIA", "HUGO IBARRA"),
    (2137, "CNJ", None, None, "tecate", "CARMEN DIAZ"),
    (718, "FORD F-600 GRUA", "1FDTEST0000000007", "ZKX0001", "TALLER", "JORGE KURI"),
]
# BASE DE DATOS: (titulo, [listas]); cada lista es (subtitulo o None, renglones)
# y cada renglon (unidad, status[, nota]). Una columna puede traer varias listas.
BASE = [
    ("ANA LOPEZ RUIZ (REPARTO)", [
        (None, [(2154, "LISTO"), (2155, "LISTO"), (2400, "LISTO"), (2201, "PENDIENTE")]),
        ("GLORIA HERRERA", [(1167, "LISTO")])]),
    ("IRMA JUAREZ (ESTACIONARIO)", [
        (None, [(2311, "LISTO"), ("BG-740", "LISTO"), ("BG-741", "LISTO")])]),
    ("DANIEL ESPINO (REPARTO)", [
        (None, [("GB-304", "PENDIENTE"), (1489, "BAJA"), ("GB-306", "LISTO")])]),
    ("CARMEN DIAZ (REPARTO)", [
        (None, [(2137, "LISTO"), (1512, "PENDIENTE", "BAJA")])]),
    ("ING. FELIPE GOMEZ (OPERACIONES)", [
        (None, [("T-822", "LISTO"), ("BGOT-04", "PENDIENTE", "TALLER"), ("BGTO-05", "LISTO"),
                ("900R", "N/A")])]),
    ("ELENA FUENTES (ESTACIONARIO)", [(None, [("BG-728", "PENDIENTE")])]),
    ("HUGO IBARRA", [
        (None, [(1203, "LISTO")]),
        ("JORGE KURI", [(718, "LISTO")])]),
]
# Bruno trae dos listas pegadas (estacionario y reparto), como dos supervisores
# del real: la nota de la primera caeria en la segunda.
BASE_DOBLE = ("BRUNO CASTRO", [
    ("ESTACIONARIO", [("BG-733", "LISTO"), ("BG-834", "LISTO"), ("BG-829", "PENDIENTE")]),
    ("REPARTO", [(2310, "LISTO"), ("BG-813", "LISTO")])])


def _hoja_base(wb, extra=()):
    b = wb.create_sheet("BASE DE DATOS")
    col = 1
    for titulo, listas in BASE:
        b.cell(1, col, titulo)
        fila = 3
        for sub, renglones in listas:
            if sub:
                b.cell(fila, col, sub)
                fila += 1
            b.cell(fila, col, "UNIDAD "), b.cell(fila, col + 1, "STATUS")
            for k, r in enumerate(renglones, start=fila + 1):
                for d, v in enumerate(r):
                    b.cell(k, col + d, v)
            fila += len(renglones) + 1
        col += 3
    titulo, listas = BASE_DOBLE
    b.cell(1, col, titulo)
    for d, (sub, renglones) in enumerate(listas):
        c = col + 2 * d
        b.cell(3, c, sub)
        b.cell(4, c, "UNIDAD "), b.cell(4, c + 1, "STATUS")
        for k, r in enumerate(renglones, start=5):
            for e, v in enumerate(r):
                b.cell(k, c + e, v)
    # Bloques de mas, a la derecha, como los pondria Logistica: (titulo, renglones).
    col += 5
    for titulo, renglones in extra:
        b.cell(1, col, titulo)
        b.cell(3, col, "UNIDAD "), b.cell(3, col + 1, "STATUS")
        for k, r in enumerate(renglones, start=4):
            for e, v in enumerate(r):
                b.cell(k, col + e, v)
        col += 3


def libro(carpeta, flotilla=FLOTILLA, gps=GPS, permisos=PERMISOS, sin=(), base_extra=(),
          permiso_en="B", permiso_dice="PERMISO ", con_canal=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PORCENTAJE"
    ws["C2"] = "TOTALES PROYECTO GPS"
    ws["C21"], ws["D21"], ws["G21"] = "ZONA", "SUPERVISOR", "TOTAL UNIDADES"
    ws[f"{permiso_en}22"] = permiso_dice
    col_perm = openpyxl.utils.column_index_from_string(permiso_en)
    for i, (canal, perm, zona, sup) in enumerate(permisos, start=23):
        ws.cell(i, 1, canal if con_canal else None), ws.cell(i, col_perm, perm)
        ws.cell(i, 3, zona), ws.cell(i, 4, sup)
    ws.cell(23 + len(permisos), 6, "SUB TOTAL")
    f = wb.create_sheet("FLOTILLA")
    f.append([None] * 6)
    f.append(["CANAL", "SUCURSAL", "SUPERVSOR", "UNIDAD", "STATUS", "STATUS"])
    for fila in flotilla:
        f.append(list(fila))
    g = wb.create_sheet("GPS LISTADO")
    g.append([None] * 8 + ["TOTAL"])
    g.append([None])
    g.append([None, "No.", "NÚMERO\nCELULAR", "MARCA, MODELO\n", "IMEIS", "UNIDAD", "MODELO",
              "VIN", "PLACAS", "INSTALACION", "CAJA", "No. GUIA", "CANAL", "SUPERVISOR"])
    g.append([None])
    for i, (u, mod, vin, pl, can, sup) in enumerate(gps, start=1):
        g.append([i, i + 0.1, 6630000000 + i, "TOPFLYTECH", "8696160606", u, mod, vin, pl,
                  datetime.datetime(2026, 2, 2), i, "N/A", can, sup])
    _hoja_base(wb, base_extra)
    for h in sin:
        del wb[h]
    ruta = os.path.join(carpeta, padron.ARCHIVO)
    wb.save(ruta)
    return ruta


def base():
    eng = create_engine("sqlite://", connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng)
    db = sessionmaker(bind=eng, autoflush=False)()
    for clave, nombre, central, _t, lat, lng in seed.PLANTAS:
        pl = m.Planta(clave=clave, nombre=nombre, es_central=central, latitud=lat, longitud=lng)
        db.add(pl)
        db.flush()
        db.add(m.Taller(planta_id=pl.id, nombre=nombre, tipo="CENTRAL" if central else "SATELITE"))
    tipos = {n: m.TipoUnidad(nombre=n) for n in ("pipa", "reparto", "utilitario")}
    db.add_all(tipos.values())
    db.flush()
    t = {x.nombre: x for x in db.query(m.Taller)}

    def unidad(num, tipo, taller, activo=True, **kw):
        u = m.Unidad(num_economico=num, tipo_unidad_id=tipos[tipo].id,
                     taller_asignado_id=t[taller].id if taller else None, activo=activo, **kw)
        db.add(u)
        return u
    # Como en produccion: la P de las pipas, una inactiva que vuelve, una con
    # placas y VIN de Logistica, utilitarios que el padron ya no trae.
    unidad("2154", "reparto", "Alamos", placas="BG0001A")
    unidad("2155", "reparto", "Alamos")
    unidad("2400", "reparto", "Alamos")
    unidad("2201", "reparto", "Alamos")              # espera su GPS: solo en BASE DE DATOS
    unidad("2310", "reparto", "Alamos", vin="LS1TEST0000000002")
    unidad("2311", "pipa", "Alamos")
    unidad("BG813", "reparto", "Rosarito")
    unidad("BG733P", "reparto", "Alamos")
    unidad("BG740P", "pipa", "Alamos")
    unidad("BG741P", "pipa", "Alamos")
    unidad("BG728P", "pipa", "Valle Redondo")
    unidad("BG304", "reparto", "Valle Redondo")
    unidad("304", "utilitario", "Alamos")            # Z Gas: otra unidad que la BG-304
    unidad("BG834", "pipa", "Rosarito", activo=False)
    unidad("1167", "utilitario", None)
    unidad("1203", "utilitario", "Alamos")
    unidad("2137", "reparto", "Tecate")
    unidad("1489", "reparto", "Valle Redondo")      # Logistica la marco BAJA
    unidad("1512", "reparto", "Tecate", activo=False)  # BAJA en la nota: no vuelve
    unidad("GRUA 718", "utilitario", "Alamos")      # Logistica la escribe '718'
    unidad("BGT-04", "utilitario", None)             # Logistica la escribe 'BGOT-04'
    unidad("PLAT 900", "utilitario", None)           # se parece al '900R' del padron
    unidad("1017", "utilitario", None)                # operacion de planta: sale del padron
    unidad("BG109", "utilitario", "Alamos")
    db.commit()
    return db


def con_libro(**kw):
    carpeta = tempfile.mkdtemp()
    try:
        return padron.leer(libro(carpeta, **kw))
    finally:
        shutil.rmtree(carpeta, ignore_errors=True)


def aplicar(db, p):
    original = padron.MINIMO_DE_UNIDADES
    padron.MINIMO_DE_UNIDADES = 1          # el libro de prueba es chico a proposito
    try:
        return padron.aplicar(db, p, hoy=HOY)
    finally:
        padron.MINIMO_DE_UNIDADES = original


def u(db, num):
    return db.query(m.Unidad).filter(m.Unidad.num_economico == num).one()


# ------------------------------------------------------------- leer ------- #
@caso("la planta es la que digan dos de las tres hojas, y la diferencia se reporta")
def _():
    p = con_libro()["unidades"]
    # FLOTILLA vieja: GPS LISTADO y BASE DE DATOS dicen Alamos.
    assert p["2311"]["planta"] == "ALAMOS", p["2311"]
    assert any("ROSARITO en FLOTILLA" in x for x in p["2311"]["a_revisar"]), p["2311"]
    # El supervisor de GPS LISTADO es el que se equivoca: FLOTILLA y BASE dicen Alamos.
    assert p["2400"]["planta"] == "ALAMOS", p["2400"]
    assert p["2310"]["planta"] == "ROSARITO" and not p["2310"]["a_revisar"]
    assert p["BG834"]["planta"] == "ROSARITO"        # solo en GPS: su supervisor y su canal
    assert p["2137"]["planta"] == "TECATE"            # el canal trae la planta
    assert p["1167"]["planta"] == "ALAMOS" and p["1167"]["canal"] == "AREA COMERCIAL"
    assert p["1203"]["planta"] is None                # 20 de Noviembre no es planta del sistema


@caso("el canal es el de FLOTILLA: la BG-813 es de reparto aunque GPS diga estacionario")
def _():
    p = con_libro()["unidades"]
    assert p["BG813"]["canal"] == "REPARTO" and p["BG813"]["permiso"] == P_VALLE
    assert any("ESTACIONARIO en GPS" in x for x in p["BG813"]["a_revisar"])


@caso("el permiso va por canal Y planta: Rosarito reparto y Rosarito estacionario son distintos")
def _():
    p = con_libro()["unidades"]
    assert p["2310"]["permiso"] == P_VALLE            # reparto Rosarito
    assert p["BG733P"]["permiso"] == P_ALAMOS         # estacionario Rosarito
    assert p["BG834"]["permiso"] == P_ALAMOS
    assert p["T822"]["permiso"] == P_VALLE            # 'OPERACIONES' = 'OPERACIÓN'
    assert p["2154"]["permiso"] == P_ALAMOS
    assert p["1167"]["permiso"] is None               # area comercial: sin permiso


@caso("BASE DE DATOS: las que esperan su GPS entran con su lista; las marcadas BAJA, no")
def _():
    r = con_libro()
    p = r["unidades"]
    assert p["2201"]["hojas"] == ["BASE DE DATOS"]
    assert (p["2201"]["canal"], p["2201"]["planta"], p["2201"]["permiso"]) == \
        ("REPARTO", "ALAMOS", P_ALAMOS)
    # La pipa nueva de la lista doble (la nota de la izquierda es la otra lista).
    assert (p["BG829"]["canal"], p["BG829"]["planta"], p["BG829"]["permiso"]) == \
        ("ESTACIONARIO", "ROSARITO", P_ALAMOS)
    assert p["1167"]["canal"] == "AREA COMERCIAL"     # la segunda lista de la columna
    assert "1489" not in p and "1512" not in p        # BAJA en el status y en la nota
    assert sum("BAJA" in a for a in r["avisos"]) == 2, r["avisos"]
    assert "2310" in p and "BG813" in p and "BG733P" in p


@caso("la misma unidad escrita distinto se junta: BG-728/BG-728P y las erratas GB y BGTO")
def _():
    p = con_libro()["unidades"]
    assert "BG728P" not in p and p["BG728"]["planta"] == "VALLEREDONDO"
    assert p["BG728"]["vin"] == "3BKTEST0000000003"   # el VIN del GPS no se pierde al juntar
    assert "GB304" not in p and p["BG304"]["planta"] == "VALLEREDONDO"
    assert "304" not in p                              # la BG-304 NO es la 304
    assert "BGT04" in p and "BGT05" in p and "BGOT04" not in p


@caso("un libro sin sus hojas no se lee; sin BASE DE DATOS solo se avisa")
def _():
    try:
        con_libro(sin=("FLOTILLA",))
    except RuntimeError as e:
        assert "FLOTILLA" in str(e)
    else:
        raise AssertionError("leyo un libro sin FLOTILLA")
    r = con_libro(sin=("BASE DE DATOS",))
    assert "2201" not in r["unidades"] and any("BASE DE DATOS" in a for a in r["avisos"])


# ----------------------------------------------------------- aplicar ------ #
@caso("las del padron quedan activas con su planta, su tipo y su permiso; las demas, de baja")
def _():
    db = base()
    r = aplicar(db, con_libro())
    t = {x.nombre: x.id for x in db.query(m.Taller)}
    assert u(db, "2310").taller_asignado_id == t["Rosarito"]
    assert u(db, "2311").taller_asignado_id == t["Alamos"]        # no se la llevo FLOTILLA
    assert u(db, "BG733P").tipo.nombre == "pipa" and u(db, "BG733P").permiso_hidrocarburos == P_ALAMOS
    assert u(db, "BG834").activo and "BG834" in r["detalle"]["activadas"]
    assert u(db, "2137").permiso_hidrocarburos == P_TECATE
    assert u(db, "1167").permiso_hidrocarburos is None and u(db, "1167").activo
    assert u(db, "2201").activo and u(db, "2201").permiso_hidrocarburos == P_ALAMOS
    # El 1203 no tiene planta del sistema: conserva la que tenia.
    assert u(db, "1203").taller_asignado_id == t["Alamos"]
    for num in ("1017", "BG109", "304", "1489"):
        assert not u(db, num).activo and u(db, num).fecha_baja == HOY, num
    assert sorted(r["detalle"]["dadas_de_baja"]) == ["1017", "1489", "304", "BG109"], r
    assert not u(db, "1512").activo                   # marcada BAJA: no se reactiva
    # Las que no estaban se crean, con el tipo de su canal y bien escritas.
    assert r["detalle"]["creadas"] == ["BG-306", "BG-829", "BGT-05", "T-822"], r
    assert u(db, "BG-829").tipo.nombre == "pipa" and u(db, "BG-829").origen == "padron"
    t822 = u(db, "T-822")
    assert t822.activo and t822.origen == "padron" and t822.tipo.nombre == "utilitario"


@caso("no duplica las que la base guarda con otro nombre: GRUA 718, BGT-04, PLAT 900")
def _():
    db = base()
    r = aplicar(db, con_libro())
    grua = u(db, "GRUA 718")
    assert grua.activo and grua.vin == "1FDTEST0000000007"     # el '718' del GPS es ella
    assert u(db, "BGT-04").activo                              # el 'BGOT-04' de Logistica
    assert not db.query(m.Unidad).filter(m.Unidad.num_economico.in_(
        ("718", "BGOT-04", "900R"))).count()
    # '900R' se parece a 'PLAT 900': ni se crea ni se da de baja la parecida.
    assert u(db, "PLAT 900").activo
    assert any(x.startswith("900R:") and "PLAT 900" in x for x in r["detalle"]["a_revisar"]), r


@caso("una nueva con errata (GB-306, BGTO-05) se crea bien escrita y no se vuelve a crear")
def _():
    db = base()
    aplicar(db, con_libro())
    assert u(db, "BG-306").activo and u(db, "BGT-05").activo
    r = aplicar(db, con_libro())                   # antes tronaba: UNIQUE num_economico
    assert r["creadas"] == 0, r


@caso("la base y el padron se cruzan con la misma clave: la 'GB-308' de la base es la BG-308")
def _():
    db = base()
    tipo = db.query(m.TipoUnidad).filter_by(nombre="reparto").one()
    db.add(m.Unidad(num_economico="GB-308", tipo_unidad_id=tipo.id, activo=True))
    db.commit()
    r = aplicar(db, con_libro(gps=GPS + [("BG-308", "CNJ", None, None, "REPARTO",
                                          "DANIEL ESPINO")]))
    assert "BG-308" not in r["detalle"]["creadas"] and "GB-308" not in r["detalle"]["dadas_de_baja"]
    assert u(db, "GB-308").activo and u(db, "GB-308").permiso_hidrocarburos == P_VALLE

@caso("placas y VIN: solo los que faltan; lo distinto se reporta y no se pisa")
def _():
    db = base()
    r = aplicar(db, con_libro())
    assert u(db, "2154").placas == "BG0001A"                   # igual, sin guiones
    assert u(db, "2154").vin == "LS1TEST0000000001"            # faltaba: se pone
    assert u(db, "2310").vin == "LS1TEST0000000002"            # distinto: NO se pisa
    assert any(x.startswith("2310: vin") for x in r["detalle"]["a_revisar"]), r
    assert u(db, "2310").placas == "BG0002A"
    assert u(db, "BG728P").vin == "3BKTEST0000000003"


@caso("correrlo dos veces no cambia nada")
def _():
    db = base()
    aplicar(db, con_libro())
    antes = [(x.num_economico, x.activo, x.taller_asignado_id, x.tipo_unidad_id,
              x.permiso_hidrocarburos, x.placas, x.vin) for x in db.query(m.Unidad).order_by(m.Unidad.id)]
    r = aplicar(db, con_libro())
    despues = [(x.num_economico, x.activo, x.taller_asignado_id, x.tipo_unidad_id,
                x.permiso_hidrocarburos, x.placas, x.vin) for x in db.query(m.Unidad).order_by(m.Unidad.id)]
    assert antes == despues
    assert not (r["creadas"] or r["activadas"] or r["dadas_de_baja"]
                or r["planta_cambiada"] or r["permiso_puesto"] or r["tipo_corregido"]
                or r["programas_creados"]), r


@caso("las nuevas y las que vuelven reciben su programa de preventivo; las demas no otro")
def _():
    db = base()
    tipos = {x.nombre: x.id for x in db.query(m.TipoUnidad)}
    planes = {t: m.PlanMantenimiento(nombre=f"prev {t}", tipo_unidad_id=tipos[t],
                                     periodicidad_dias=d, activo=True)
              for t, d in (("pipa", 60), ("reparto", 90))}
    db.add_all(planes.values())
    db.flush()
    # La flota que ya estaba tiene su programa abierto, como en produccion.
    for x in db.query(m.Unidad).filter(m.Unidad.activo.is_(True)):
        plan = next((p for p in planes.values() if p.tipo_unidad_id == x.tipo_unidad_id), None)
        if plan:
            db.add(m.ProgramaMantenimiento(unidad_id=x.id, plan_id=plan.id,
                                           fecha_limite=HOY, estado="pendiente"))
    db.commit()
    r = aplicar(db, con_libro())

    def programas(num):
        return (db.query(m.ProgramaMantenimiento)
                .filter(m.ProgramaMantenimiento.unidad_id == u(db, num).id).all())
    # BG834 vuelve y BG-829 es nueva (pipas); BG-306 es nueva de reparto.
    for num in ("BG834", "BG-829", "BG-306"):
        assert len(programas(num)) == 1 and programas(num)[0].fecha_limite > HOY, num
    assert len(programas("2154")) == 1                 # ya tenia: no se le crea otro
    assert not programas("T-822")                      # utilitario: sin plan
    assert r["programas_creados"] == 3, r
    assert aplicar(db, con_libro())["programas_creados"] == 0


@caso("PORCENTAJE: el permiso y el canal se reconocen por lo que traen, no por su lugar")
def _():
    # Encabezado reescrito y el permiso en otra columna: antes se perdian los
    # supervisores y FLOTILLA (vieja) volvia a mandar la planta.
    r = con_libro(permiso_en="E", permiso_dice="PERMISO CRE")
    p = r["unidades"]
    assert p["2310"]["permiso"] == P_VALLE and p["BG733P"]["permiso"] == P_ALAMOS
    assert p["2311"]["planta"] == "ALAMOS", p["2311"]       # el supervisor sigue contando
    assert not any("permiso" in a.lower() for a in r["avisos"] if "BAJA" not in a), r["avisos"]


@caso("PORCENTAJE sin el canal de sus supervisores: no se aplica a medias, se detiene")
def _():
    try:
        con_libro(con_canal=False)
    except RuntimeError as e:
        assert "supervisores" in str(e) and "No se" in str(e), e
    else:
        raise AssertionError("leyo un PORCENTAJE sin canal y siguio")


@caso("una lista de BASE DE DATOS sin ninguna unidad con GPS entra con lo que dice su titulo")
def _():
    r = con_libro(base_extra=[
        ("BRUNO CASTRO (ESTACIONARIO)", [("BG-950", "PENDIENTE"), ("BG-951", "PENDIENTE")]),
        ("PERSONA NUEVA (REPARTO)", [("BG-960", "PENDIENTE")])])
    p = r["unidades"]
    assert (p["BG950"]["canal"], p["BG950"]["planta"], p["BG950"]["permiso"]) ==         ("ESTACIONARIO", "ROSARITO", P_ALAMOS), p["BG950"]
    assert "BG951" in p
    # Un supervisor que PORCENTAJE no conoce: el canal si, la planta no (se avisa).
    assert (p["BG960"]["canal"], p["BG960"]["planta"]) == ("REPARTO", None), p["BG960"]
    sin_gps = [a for a in r["avisos"] if "ninguna unidad con GPS" in a]
    assert len(sin_gps) == 2 and any("sin planta" in a for a in sin_gps), r["avisos"]


@caso("un permiso corregido a mano en la bitacora no se pisa en la siguiente sincronizacion")
def _():
    db = base()
    aplicar(db, con_libro())
    bg813 = u(db, "BG813")
    assert bg813.permiso_hidrocarburos == P_VALLE
    admin = m.Usuario(nombre="Admin", apellidos="Prueba", email="admin@prueba.test",
                      password_hash="x")
    rol = m.Rol(nombre="administrador")
    db.add_all([admin, rol])
    db.flush()
    db.add(m.UsuarioRol(usuario_id=admin.id, rol_id=rol.id))
    db.commit()
    # El administrador lo corrige (la BG-813 tiene el canal en duda).
    bc.fijar_identificacion(bg813.id, bc.IdentificacionUnidadIn(permiso_hidrocarburos=P_ALAMOS),
                            usuario=admin, db=db)
    asientos = db.query(m.AsientoBitacora).filter(m.AsientoBitacora.unidad_id == bg813.id).count()
    r = aplicar(db, con_libro())
    assert u(db, "BG813").permiso_hidrocarburos == P_ALAMOS
    assert r["permiso_puesto"] == 0, r
    assert any(x.startswith("BG813: permiso") and "a mano" in x for x in r["detalle"]["a_revisar"])
    assert db.query(m.AsientoBitacora).filter(
        m.AsientoBitacora.unidad_id == bg813.id).count() == asientos   # sin idas y vueltas
    # Otro dato corregido a mano (el auxiliar) no congela el permiso.
    otra = u(db, "2154")
    otra.permiso_hidrocarburos = None
    db.commit()
    bc.fijar_identificacion(otra.id, bc.IdentificacionUnidadIn(personal_auxiliar="Ayudante"),
                            usuario=admin, db=db)
    aplicar(db, con_libro())
    assert u(db, "2154").permiso_hidrocarburos == P_ALAMOS


@caso("el permiso nuevo se asienta en el libro de la NOM-030 de la unidad que ya tiene uno")
def _():
    db = base()
    con = u(db, "2154")
    bs.asentar(db, unidad_id=con.id, tipo="apertura", descripcion="Apertura de prueba",
               origen="formato_mantenimiento")
    db.commit()
    aplicar(db, con_libro())
    a = (db.query(m.AsientoBitacora)
         .filter(m.AsientoBitacora.unidad_id == con.id, m.AsientoBitacora.tipo == "identificacion")
         .one())
    assert P_ALAMOS in a.descripcion and "padrón de Logística" in a.descripcion, a.descripcion
    sin = u(db, "2310")
    assert not db.query(m.AsientoBitacora).filter(m.AsientoBitacora.unidad_id == sin.id).count()
    v = bs.verificar_libro(db, con.id, db.query(m.AsientoBitacora)
                           .filter(m.AsientoBitacora.unidad_id == con.id)
                           .order_by(m.AsientoBitacora.numero).all())
    assert v["integra"], v


@caso("un padron que trae muy pocas unidades no se aplica: no da de baja la flota")
def _():
    db = base()
    p = con_libro()
    try:
        padron.aplicar(db, p, hoy=HOY)       # con el minimo real
    except RuntimeError as e:
        assert "no se" in str(e).lower() and "aplic" in str(e).lower(), e
    else:
        raise AssertionError("aplico un padron de unas cuantas unidades")
    db.rollback()
    assert u(db, "1017").activo and u(db, "BG109").activo


@caso("el buscador del formato encuentra tambien las de baja, marcadas")
def _():
    db = base()
    aplicar(db, con_libro())
    r = ac.buscar_unidades(q="304", limite=10, usuario=None, db=db)
    assert [(x["num_economico"], x["activo"]) for x in r] == [("304", False), ("BG304", True)], r


@caso("el buscador pone primero la exacta aunque este de baja (si no, no cabe)")
def _():
    db = base()
    tipo = db.query(m.TipoUnidad).first()
    db.add(m.Unidad(num_economico="BG-35", tipo_unidad_id=tipo.id, activo=False))
    db.add_all(m.Unidad(num_economico=f"BG35{i}P", tipo_unidad_id=tipo.id, activo=True)
               for i in range(1, 10))
    db.commit()
    r = ac.buscar_unidades(q="BG-35", limite=8, usuario=None, db=db)
    assert r[0]["num_economico"] == "BG-35" and not r[0]["activo"], r


@caso("la sincronizacion corre el padron al final y ya no lee UNIDADES BAJA GAS")
def _():
    pasos = [p[0] for p in sincronizar._pasos("/x")]
    assert pasos[-1] == "padron" and "flota" not in pasos and "catalogo" not in pasos, pasos
    assert sincronizar.ARCHIVOS["padron"] == ["CONTROL GPS.xlsx"]
    assert cs.TIPOS["unidades"]["archivo"] == "CONTROL GPS.xlsx"
    assert cs.TIPOS["unidades"]["pasos"] == ("padron",)


@caso("el archivo real, si esta en datos/: cientos de unidades y los tres permisos")
def _():
    ruta = os.path.join(os.path.dirname(BACKEND), "datos", padron.ARCHIVO)
    if not os.path.exists(ruta):
        return "se salta: no esta datos/CONTROL GPS.xlsx"
    p = padron.leer(ruta)
    assert len(p["unidades"]) >= padron.MINIMO_DE_UNIDADES
    assert {P_ALAMOS, P_VALLE, P_TECATE} <= set(p["permisos"].values())
    solo_base = [x for x in p["unidades"].values() if x["hojas"] == ["BASE DE DATOS"]]
    assert solo_base and all(x["planta"] or x["canal"] == "TESORERIA" for x in solo_base)
    return f"{len(p['unidades'])} unidades, {len(solo_base)} esperando su GPS"


if __name__ == "__main__":
    fallas = 0
    for nombre, fn in CASOS:
        try:
            extra = fn()
            print(f"  ok    {nombre}" + (f"  ({extra})" if extra else ""))
        except Exception:
            fallas += 1
            print(f"  FALLA {nombre}")
            traceback.print_exc()
    print(f"\n{len(CASOS) - fallas} de {len(CASOS)} pasaron")
    sys.exit(1 if fallas else 0)
